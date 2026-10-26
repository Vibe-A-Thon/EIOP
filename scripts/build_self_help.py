"""Generate the offline EOIP help site and matching printable manuals.

Content is intentionally tied to the current apps/web UI and documented local
deployment. Conditional services are identified as such, not shown as proven.
"""

from __future__ import annotations

from html import escape
from pathlib import Path
import json

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "self-help-manual"
SITE = OUT / "site"
DOCS = OUT / "docs"
SCREEN = SITE / "assets" / "connection-screen.png"
DASHBOARD = ROOT / ".audit-scratch" / "dashboard-desktop.png"


PAGES = [
    dict(slug="index", title="EOIP Help Center", category="Start here", audience="Everyone",
         summary="Practical instructions for the current EOIP web application, local deployment, graph exploration and operations.",
         status="Documentation snapshot - 20 October 2026", image="connection-screen.png",
         sections=[
             ("Choose your path", "The default web application is apps/web. Connect with a deployment access token, then open the Knowledge explorer, Ontology studio, Cognitive dashboards, Administration, Guided demo or Appearance from the workspace navigation. A separate frontend/ prototype exists and is not the basis of these instructions.", []),
             ("What is verified here", "Navigation labels, controls, API calls and local start commands were checked against this repository. The unauthenticated connection screen was captured from the running Vite application. Authenticated workflows need a running API, configured backend and permission to the relevant endpoints; they were not certified by the documentation build.", []),
             ("Start with a safe local environment", "For an integrated persistent environment, use the container helper from the repository root. For a developer-only session, use the memory graph profile. Memory graph data disappears when the API process exits; it is not a substitute for a persistent deployment.", ["Read Connect and first session for exact launch commands and token handling.", "Use the Knowledge explorer guide to inspect a bounded result before trying changes.", "Administrators should read data operations and security before importing a real source."]),
             ("Product status", "The separate project status dashboard is an offline audit snapshot. It explicitly distinguishes code and test leads from accepted runtime behavior. Do not treat a planned portal, a catalog entry or a diagram as an installed feature.", []),
         ]),
    dict(slug="connect", title="Connect and first session", category="Start here", audience="New users and developers",
         summary="Start a local profile, check readiness, connect with a token and find the workspace.",
         status="Available screen; backend required for sign in", image="connection-screen.png",
         sections=[
             ("Before you begin", "Use Python 3.11 or later and Node.js 22 for host development. The integrated profile requires a supported OCI container runtime with Compose. Obtain the token from the deployment owner; do not put it in chat, screenshots, source files or support tickets.", []),
             ("Option A - integrated containers", "From the repository root in PowerShell, run .\\scripts\\container.ps1 detect, then .\\scripts\\container.ps1 up. On Linux or macOS use ./scripts/container.sh detect and ./scripts/container.sh up. The helper creates a private deploy/.env on first use, with generated EOIP_API_TOKEN and HUGEGRAPH_PASSWORD. Wait for health, then open http://localhost:8080/app. The API documentation is at http://localhost:8000/docs.", ["Run the detect command and resolve any missing runtime before starting.", "Run up and wait for service health instead of assuming an open port means ready.", "Read EOIP_API_TOKEN from deploy/.env locally without exposing it in a command transcript.", "Open the web explorer, enter the token, and choose Connect."]),
             ("Option B - host-only development", "Create a virtual environment and install the development extra with python -m pip install -e \".[dev]\". Set EOIP_API_TOKEN to a disposable local value of at least 24 characters, EOIP_GRAPH_BACKEND=memory, and EOIP_DATA_DIR to a private local directory. Start python -m uvicorn apps.api.main:create_app --factory --host 127.0.0.1 --port 8000 --reload. In a second shell, run npm --prefix apps/web ci, set EOIP_DEV_API_URL=http://127.0.0.1:8000, then run npm --prefix apps/web run dev. Use the Vite URL shown by the terminal.", ["Keep the API and Vite shells running while you use the web UI.", "Check http://127.0.0.1:8000/health/ready before connecting.", "Do not start a separate memory-backed worker and expect it to share graph state with the API."]),
             ("Connect to EOIP", "The connection page displays local service readiness and an Access token field. The show/hide control changes only field visibility. The token is held in browser memory until disconnect or reload. Connect calls administrative summary endpoints, so a token can still be denied by policy even when /health/ready reports ready.", ["Confirm the local service indicator says ready. If unavailable, inspect the API/proxy before entering the token.", "Enter the deployment token and select Connect.", "Wait for the workspace shell. If an error toast appears, record its correlation ID but never the token.", "Use Disconnect when finished; reload also clears the in-memory token."]),
             ("Find your way", "The workspace navigation contains Knowledge explorer, Cognitive Engineering, Cognitive Dashboards, Ontology studio, Administration, Guided demo and Appearance. Roles and backend configuration can restrict actions. A visible menu item is not proof that its backing service is configured.", []),
         ]),
    dict(slug="manuals", title="Printable manuals", category="Start here", audience="Everyone",
         summary="Download the same guidance in PDF or editable Word format for training and offline reference.",
         status="Five guide sets available", image=None,
         sections=[
             ("Choose a manual", "The quick start covers first connection. The user guide covers graph exploration and studio work. The administrator guide covers sources, validation and access. Operations and troubleshooting covers diagnosis and recovery. Documentation and product status explains the separate project audit dashboard.", []),
             ("Keep a current copy", "The website and manuals are generated from the same content source. Re-run python scripts/build_self_help.py after documentation edits, then render and inspect the PDF and DOCX outputs. Screenshots reflect a disposable development profile and must not be treated as production acceptance evidence.", []),
         ]),
    dict(slug="explore", title="Explore the knowledge graph", category="Workflows", audience="Analysts and engineers",
         summary="Search, inspect entities and edges, trace impact and read evidence without assuming a bounded view is the entire graph.",
         status="UI implemented; result quality depends on graph and access", image="explorer-screen.png",
         sections=[
             ("Search for a known entity", "Open Knowledge explorer. Search by an issue key, symbol, commit, file or document identifier you already know. Results are grouped by knowledge domain and retain canonical graph IDs. An empty result can mean no matching data, an access restriction or an unavailable backend.", ["Enter a specific identifier in the search field and submit.", "Read the match count and any partial-result warning.", "Choose a result in its domain group to select the canonical entity."]),
             ("Inspect an entity", "The right inspector shows Overview and Evidence tabs. Overview lists properties; Evidence shows connected provenance where available. Property values and source references are data from the graph, not an independent fact check. Filter a large property list using the inspector filter.", ["Select an entity in results, the entity strip or graph canvas.", "Inspect its type and canonical ID before sharing a link or using it in a path request.", "Open Evidence and check source references; note when no direct evidence is returned."]),
             ("Read the relationship map", "The graph lens shows a bounded subgraph. Filter by domain, relationship class and deterministic or inferred trust state. A 'partial' warning means more context may exist. Select an edge to inspect direction, type, properties and evidence key. Thicker links indicate attached evidence, not necessarily validated truth.", ["Use filters to reduce clutter before interpreting connections.", "Select a relationship and compare from and to IDs.", "Expand a selected entity only when you need more context; results remain policy bounded."]),
             ("Trace downstream impact", "Select an anchor entity, then use Explore impact. The current UI describes a four-hop bound. Run TraceImpact invokes a deterministic read-only agent. Read the per-node rationale where returned; do not interpret absence of a path as evidence that no dependency exists outside the permitted scope.", ["Select an entity in Graph and switch to the Impact lens.", "Run Explore impact, then inspect the returned path depth and evidence.", "For two known IDs, enter both in Trace requirement to implementation and choose Find path."]),
             ("Other lenses", "Architecture shows the domains and relationship classes in the current bounded result, not a full inventory. The lens library distinguishes API-backed choices from disabled planned lenses. A disabled lens should remain disabled until its backend is delivered and qualified.", []),
         ]),
    dict(slug="studio", title="Query and ontology studio", category="Workflows", audience="Analysts and data stewards",
         summary="Use the studio with known-answer queries and inspect schema before interpreting a result.",
         status="Conditional - backend dialect and role determine availability", image="studio-screen.png",
         sections=[
             ("Open the studio", "Choose Ontology studio from the workspace navigation. Its schema and query tools communicate with the configured API. A development memory graph may be empty; load only an approved demo or source dataset for meaningful results.", ["Connect to a ready API with an authorized token.", "Open Ontology studio and inspect available schema or sample content.", "Confirm the backend and ontology version before writing a query."]),
             ("Run a controlled query", "Start with a small known-answer question against an entity you are allowed to view. Use the studio's supported query form and limits; do not paste secrets or an unrestricted data export into a query. Check the returned count, warnings, provenance and any truncation. Compare one result with Knowledge explorer before relying on it.", ["Choose the supported query mode for the active backend.", "Use an exact identifier or narrow filter and a bounded limit.", "Execute, inspect errors and compare results with a known entity.", "Refine one condition at a time and retain the final query for review."]),
             ("Understand ontology versions", "The administration view can load graph schema and inspect a JSON-LD context. Schema visibility does not grant permission to change ontology. Treat a version change as a migration with compatibility and validation checks.", ["Ask an administrator which schema version is active.", "Review permitted labels and relationships before mapping data.", "Run validation after approved ingestion or ontology changes."]),
             ("When a query is unsupported", "A dialect shown in requirements or code is not automatically qualified for the active graph backend. Preserve the query text, backend name, error and correlation ID, then use the deployment's supported query path. Do not silently rewrite an authorization failure as an empty result.", []),
         ]),
    dict(slug="data-admin", title="Ingest and validate data", category="Administration", audience="Administrators and data owners",
         summary="Operate demo loading, file import, source registry, dry runs, jobs and validation in the current Administration UI.",
         status="Conditional - configured connectors and worker required", image="sources-screen.png",
         sections=[
             ("Start with a disposable dataset", "Administration exposes a Demo dataset action for synthetic Jira, Git, code, pipeline and documentation evidence. Use a non-production environment and confirm that demo loading is permitted. For a production deployment set EOIP_LOAD_DEMO=false as documented by the repository. In the tested memory profile, the service panel misleadingly said HugeGraph READY; verify the actual backend configuration separately.", ["Open Administration and review system and graph status.", "In a disposable environment, choose Load dataset.", "Refresh counts and inspect the graph inventory before querying."]),
             ("Import a JSON file", "The UI has Jira JSON and GitLab metadata file import controls, and a JSON graph-ingestion area. Validate the source format and classification with the data owner first. Use a small, approved sample to check record-level errors before larger imports. A button indicates a request path; verify the resulting graph, error list and job state.", ["Choose the correct import source and select the approved JSON file.", "Run the import and wait for the result or job ID.", "Review accepted, rejected and errored records, not only a success banner.", "Search for one imported canonical entity and inspect provenance."]),
             ("Operate registered sources", "The source registry exposes baseline, discovery, dry run, sync, reconciliation and replay actions. Named connectors such as Confluence, ServiceNow, Rally, Grafana, Datadog, Kibana, qTest, Figma, API Portal and GitHub require external configuration; do not assume credentials or access are present. Confluence and ServiceNow incremental options include checkpoint and Dry run only.", ["Load source control and choose an already registered, authorized source.", "Review its endpoint, owner, scope and status without revealing credentials.", "Run discovery or dry run first; record counts and rejected records.", "Queue sync only after approval and monitor its durable job.", "Reconcile source and graph counts, then inspect sampled entities and evidence."]),
             ("Validate and recover", "Run validation from Knowledge quality after a data change. Check evidence coverage, provenance coverage, schema compliance, completeness, freshness, orphan and duplicate rates plus individual findings. The Dead-letter queue lists rejection context without captured payloads. Replay only after the underlying cause is corrected and the data owner approves it.", ["Choose Run validation and record the checked entity and relationship counts.", "Review each error and warning, then correct source mapping or data.", "Use Load DLQ to inspect failed records and the reason, never to extract sensitive payloads.", "Re-run a bounded dry run or replay and confirm no duplicate effects."]),
         ]),
    dict(slug="reports", title="Dashboards and project status", category="Workflows", audience="Analysts and delivery leads",
         summary="Use the application dashboards and the separate offline audit tracker without confusing their purposes.",
         status="Conditional app dashboards; offline audit snapshot available", image="dashboards-screen.png",
         sections=[
             ("Cognitive dashboards", "Choose Cognitive Dashboards in the signed-in workspace. Panels use the active API and data profile. Read panel titles, time range, source and warnings before drawing conclusions. If a panel is empty, check whether the dataset, permission and backing service exist; do not fill gaps with made-up figures.", ["Open the dashboard after connecting to a healthy API.", "Inspect the source and scope of each visible measure.", "Refresh after a known data change and compare with source counts.", "Preserve a filtered view or export only through an authorized workflow."]),
             ("Offline project audit dashboard", "EOIP-master-dashboard.html at the repository root is a separate project-delivery tracker, not the application analytics portal. Open it directly in a browser. Tabs cover overview, requirements, gaps, delivery, evidence, sources, architecture, checks and decisions. The file contains a point-in-time snapshot; rebuild with python scripts/build_unified_audit.py after source changes.", ["Open the root dashboard HTML without a server.", "Search a requirement ID and inspect both its reported state and audit state.", "Use Gaps and Evidence to see closure actions and evidence limits.", "Export filtered CSV or local decision notes; these exports do not change acceptance."]),
             ("Interpret status honestly", "The current structural audit indexed thousands of requirement IDs but independently certified none in a full runtime release gate. 'Code observed' means a referenced path exists, not that the user journey works. 'Claimed accepted' preserves a historical ledger statement, not an independently verified production outcome.", []),
         ]),
    dict(slug="security", title="Security and access", category="Administration", audience="Administrators and security reviewers",
         summary="Handle tokens, tenant scope, audit events, egress and derived artifacts as operational controls.",
         status="Controls require deployment-level verification", image="administration-screen.png",
         sections=[
             ("Protect the access token", "The primary web UI keeps the token in memory and removes an older sessionStorage credential. Reload or Disconnect clears it. This does not replace HTTPS, secret rotation, least privilege or endpoint authorization in a real deployment. Do not share credentials in screenshots or support bundles.", ["Issue a scoped token through the approved deployment process.", "Use the primary apps/web UI and a trusted origin.", "Disconnect when complete; rotate a token if exposure is suspected."]),
             ("Review audit events", "Administration has a Load audit events action. The screen aims to show safe summaries, actors, outcomes and correlation identifiers without secret payloads. Review any action that changes source registration, graph state, policy, job recovery or export access.", ["Load the latest events and identify the relevant correlation ID.", "Check actor, tenant, operation, outcome and time against the approved request.", "Escalate missing, ambiguous or cross-tenant entries; do not infer a clean audit from an empty UI list."]),
             ("Tenant and data boundaries", "The project intends one authorization authority across UI, API, worker, MCP and developer channels. Current code requires further end-to-end qualification, including derived data and downloads. Administrators should test denied cross-tenant reads and revocation on the actual deployment before trusting a release claim.", []),
             ("External services and local models", "Connectors need explicit credentials and source allowlists. Local AI and processing backends are optional. Application-level egress checks are not a network firewall; enforce no-egress at the host or orchestration layer where the data classification requires it, and verify with observed network tests.", []),
         ]),
    dict(slug="troubleshoot", title="Troubleshooting and recovery", category="Support", audience="Users, administrators and support",
         summary="Diagnose readiness, connection, empty results, stalled jobs and failed imports without leaking sensitive data.",
         status="Read-only first; recovery actions need approval", image="connection-screen.png",
         sections=[
             ("Service status unavailable", "This is the expected connection-screen state when the web frontend runs without an API proxy or the API is unhealthy. Check /health/ready on the API. In Vite host development, set EOIP_DEV_API_URL before starting Vite. In containers, inspect the compose status and service logs. Do not repeatedly submit a token until readiness is understood.", ["Record the time and URL, not the token.", "Check the API readiness endpoint and container status.", "Confirm the correct web origin and API proxy, then reload."]),
             ("Cannot connect", "Connect requests administrative stats, system and jobs. A ready health endpoint can coexist with an invalid token or denied authorization. Read the toast and correlation ID. Confirm the deployment token and scope through the administrator; never paste the token into a ticket.", ["Check readiness first.", "Record the error text and correlation ID.", "Ask the deployment owner to verify token scope or rotate a compromised token."]),
             ("No search results", "Check the spelling and canonical ID, selected filters, source ingestion status and tenant access. A bounded or partial graph view can omit other permitted context. Compare one known imported record; do not conclude that the source has no data from a single empty query.", []),
             ("Import or job stalled", "Use Administration to inspect durable queue state and dead-letter records. Distinguish pending, running, failed, cancelled and completed. Check worker health, lease behavior and storage before any replay. Re-run only with an idempotency and duplicate-effects check.", ["Capture the job ID and error summary, not sensitive payload.", "Inspect worker status and source checkpoint.", "Correct the root cause and obtain approval for replay.", "Reconcile source counts, persisted graph state and audit events afterward."]),
             ("Backups and restoration", "The integrated container helper exposes backup and restore-test actions. A completed backup label is not proof of recoverability. Use the operations runbook and a disposable environment to perform a restore drill before relying on a snapshot. Do not delete named volumes merely to reset a test without confirming the target.", []),
             ("Support bundle", "Include application version or git revision, deployment profile, UTC time, affected route, steps to reproduce, correlation/job ID, expected versus observed behavior and redacted logs. Exclude access tokens, raw source payloads, private graph exports and model prompts unless an approved secure channel exists.", []),
         ]),
]


def html_page(page, search_index):
    slug = page["slug"]
    title = escape(page["title"])
    groups = ["Start here", "Workflows", "Administration", "Support"]
    nav = "".join(
        f'<div class="nav-group"><span>{escape(group)}</span>' + "".join(
            f'<a href="{p["slug"]}.html" class="{ "current" if p["slug"] == slug else "" }" {"aria-current=page" if p["slug"] == slug else ""}>{escape(p["title"])}</a>'
            for p in PAGES if p["category"] == group) + "</div>"
        for group in groups
    )
    sections = ""
    toc = ""
    for i, (heading, description, steps) in enumerate(page["sections"], 1):
        ident = f"section-{i}"
        toc += f'<a href="#{ident}">{escape(heading)}</a>'
        steps_html = ""
        if steps:
            steps_html = '<ol class="steps">' + "".join(f'<li><span>{i}</span><p>{escape(step)}</p></li>' for i, step in enumerate(steps, 1)) + '</ol>'
        sections += f'<section id="{ident}" class="doc-section"><div class="section-index">{i:02d}</div><div><h2>{escape(heading)}</h2><p>{escape(description)}</p>{steps_html}</div></section>'
    if slug == "manuals":
        sections += '<div class="downloads">' + "".join(f'<div><strong>{escape(name)}</strong><span><a href="../docs/{file_slug}.pdf" download>PDF</a><a href="../docs/{file_slug}.docx" download>DOCX</a></span></div>' for name, file_slug, _ in GUIDES) + '</div>'
    image = ""
    if page["image"]:
        asset = page["image"]
        caption = "Actual unauthenticated EOIP connection screen, captured from apps/web with the API unavailable." if asset == "connection-screen.png" else "Actual EOIP web screen in a disposable memory-backed local development session with no source data loaded."
        image = f'<figure class="product-shot"><a href="assets/{asset}" target="_blank" rel="noopener" title="Open screenshot at full size"><img src="assets/{asset}" alt="{escape(caption)}" loading="lazy"></a><figcaption>{escape(caption)} <a href="assets/{asset}" target="_blank" rel="noopener">Open full size</a></figcaption></figure>'
    current = PAGES.index(page)
    pager = '<div class="page-turn">' + (f'<a href="{PAGES[current-1]["slug"]}.html">Previous: {escape(PAGES[current-1]["title"])}</a>' if current else '<span></span>') + (f'<a href="{PAGES[current+1]["slug"]}.html">Next: {escape(PAGES[current+1]["title"])}</a>' if current + 1 < len(PAGES) else '<span></span>') + '</div>'
    search_data = json.dumps(search_index, ensure_ascii=True).replace('<', '\\u003c')
    return f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="color-scheme" content="light"><title>{title} | EOIP Help</title><link rel="stylesheet" href="assets/help.css"></head><body>
<div class="layout"><aside class="sidebar"><a class="identity" href="index.html"><span class="mark">E</span><span><strong>EOIP</strong><small>Help center</small></span></a><button id="search-open" class="search-trigger" type="button" aria-label="Search all help pages"><span>Search documentation</span><kbd>/</kbd></button><nav aria-label="Documentation">{nav}</nav><div class="sidebar-foot"><a href="manuals.html">Printable manuals</a><a href="../../EOIP-master-dashboard.html">Project status</a><span>Local offline documentation</span></div></aside>
<div class="main-wrap"><header class="mobile-head"><button id="menu-toggle" type="button" aria-label="Open navigation">Menu</button><a href="index.html">EOIP Help</a><button id="mobile-search" type="button">Search</button></header><main class="article"><div class="eyebrow">{escape(page['category'])} / {escape(page['audience'])}</div><h1>{title}</h1><p class="lead">{escape(page['summary'])}</p><div class="meta"><span class="status">{escape(page['status'])}</span><span>Updated 20 October 2026</span></div>{image}<div class="article-body">{sections}</div>{pager}<footer>Documentation describes the current repository and observed UI. Product behavior must be verified in your deployment. <a href="../../docs/getting-started.md">Source setup guide</a></footer></main><aside class="toc"><strong>On this page</strong>{toc}<div class="toc-note">Need a paper copy?<br><a href="manuals.html">Open printable manuals</a></div></aside></div></div>
<dialog id="search-dialog" class="search-dialog"><div class="search-head"><label for="global-search">Search all help pages</label><button id="search-close" type="button" aria-label="Close search">Close</button></div><input id="global-search" type="search" autocomplete="off" placeholder="Try token, graph, import, backup..."><div id="search-results" class="search-results" aria-live="polite"></div></dialog><script>window.EOIP_SEARCH={search_data};</script><script src="assets/help.js"></script></body></html>'''


CSS = r'''*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#f8fafb;color:#19323c;font:15px/1.65 Inter,Segoe UI,Arial,sans-serif}a{color:#06677d;text-decoration:none}a:hover{text-decoration:underline}button,input{font:inherit}button{cursor:pointer}:focus-visible{outline:3px solid #e5a32c;outline-offset:2px}.layout{display:grid;grid-template-columns:258px minmax(0,1fr);min-height:100vh}.sidebar{position:sticky;top:0;height:100vh;background:#fff;border-right:1px solid #e0e8ea;padding:26px 18px;display:flex;flex-direction:column;overflow:auto}.identity{display:flex;align-items:center;gap:11px;color:#19323c;margin:0 8px 30px}.identity:hover{text-decoration:none}.mark{display:grid;place-items:center;width:39px;height:39px;border-radius:7px;background:#00768d;color:#fff;font-size:21px;font-weight:800}.identity strong,.identity small{display:block;line-height:1.15}.identity strong{font-size:17px}.identity small{font-size:11px;color:#738893;margin-top:4px}.search-trigger{border:1px solid #d7e3e8;border-radius:6px;color:#687f89;background:#f7fafb;display:flex;align-items:center;justify-content:space-between;padding:8px 10px;font-size:12px;text-align:left;margin:0 8px 25px}.search-trigger:hover{border-color:#0b8797;color:#164b56}.search-trigger kbd{border:1px solid #d3dfe3;background:#fff;border-radius:3px;padding:0 6px;color:#58707b}.nav-group{margin:0 0 23px}.nav-group span{display:block;color:#8b9ca3;text-transform:uppercase;font-size:10px;font-weight:800;margin:0 10px 7px}.nav-group a{display:block;padding:7px 10px;margin:2px 0;border-radius:5px;color:#3d5760;font-size:13px;line-height:1.3}.nav-group a:hover{background:#eef7f8;text-decoration:none}.nav-group a.current{background:#e4f3f4;color:#00677d;font-weight:750}.sidebar-foot{margin-top:auto;border-top:1px solid #e8eef0;padding:18px 10px 0;display:grid;gap:7px;font-size:12px}.sidebar-foot span{color:#91a1a8}.main-wrap{display:grid;grid-template-columns:minmax(0,900px) 200px;justify-content:center;min-width:0}.article{padding:56px 5.2vw 60px;min-width:0}.eyebrow{color:#00768d;text-transform:uppercase;font-weight:800;font-size:11px;letter-spacing:.8px}.article h1{font-size:clamp(29px,3vw,39px);line-height:1.18;margin:12px 0 12px;letter-spacing:0;color:#16313c}.lead{font-size:17px;line-height:1.6;color:#54717d;max-width:700px;margin:0}.meta{display:flex;align-items:center;gap:15px;flex-wrap:wrap;color:#869aa3;font-size:11px;margin:24px 0 32px}.status{background:#edf6f4;color:#2a756c;border:1px solid #cee5df;border-radius:4px;padding:4px 7px;font-weight:700}.product-shot{margin:0 0 42px;border:1px solid #dfe8ec;background:#fff;padding:8px;border-radius:6px;box-shadow:0 8px 28px #24404d0b}.product-shot img{width:100%;display:block;aspect-ratio:16/9;object-fit:contain;background:#f4f7f8}.product-shot figcaption{font-size:11px;color:#738792;padding:8px 9px 3px}.doc-section{display:grid;grid-template-columns:44px minmax(0,1fr);gap:17px;border-top:1px solid #e0e9ec;padding:29px 0 31px;scroll-margin-top:35px}.section-index{font-size:11px;font-weight:800;color:#0b8295;padding-top:7px}.doc-section h2{font-size:21px;line-height:1.3;letter-spacing:0;margin:0 0 11px}.doc-section p{margin:0;color:#4d6570}.steps{list-style:none;margin:20px 0 0;padding:0;display:grid;gap:11px}.steps li{display:grid;grid-template-columns:23px 1fr;gap:12px;align-items:start}.steps li span{display:grid;place-items:center;width:21px;height:21px;border-radius:50%;background:#d9eff1;color:#067185;font-size:11px;font-weight:800;margin-top:1px}.steps p{font-size:14px}.toc{position:sticky;top:0;height:100vh;padding:62px 20px 30px 0;display:flex;flex-direction:column;gap:10px;font-size:12px}.toc strong{font-size:10px;text-transform:uppercase;color:#80949d;letter-spacing:.7px;margin-bottom:3px}.toc a{color:#647e88;line-height:1.35}.toc a:hover{color:#00677d}.toc-note{border-top:1px solid #dfe9eb;padding-top:18px;margin-top:17px;color:#8b9da4}.toc-note a{color:#00677d}.page-turn{display:flex;justify-content:space-between;gap:14px;border-top:1px solid #e0e9ec;padding:21px 0;margin-top:15px;font-size:13px;font-weight:700}.page-turn a{max-width:45%}.article footer{color:#82969e;font-size:11px;border-top:1px solid #e0e9ec;padding:20px 0}.mobile-head{display:none}.search-dialog{width:min(650px,calc(100vw - 32px));max-height:min(80vh,780px);padding:0;border:1px solid #d5e3e7;border-radius:8px;box-shadow:0 24px 90px #19323c42;background:#fff}.search-dialog::backdrop{background:#122c37a8;backdrop-filter:blur(2px)}.search-head{display:flex;justify-content:space-between;align-items:center;padding:18px 19px 11px;font-size:13px;font-weight:750}.search-head button{border:0;background:none;color:#067185;padding:4px}.search-dialog input{display:block;margin:0 19px 14px;width:calc(100% - 38px);border:1px solid #bdd6dc;border-radius:5px;padding:12px 13px;outline:0}.search-results{max-height:55vh;overflow:auto;border-top:1px solid #e7eff1;padding:9px}.search-results a{display:block;padding:10px 12px;border-radius:5px}.search-results a:hover,.search-results a:focus{background:#edf7f8;text-decoration:none}.search-results strong{display:block;color:#164755;font-size:13px}.search-results small{display:block;color:#8197a0;margin:2px 0}.search-results p{margin:0;color:#5c737d;font-size:12px}.search-empty{padding:18px;color:#77909a;font-size:13px}@media(max-width:1190px){.main-wrap{grid-template-columns:minmax(0,900px)}.toc{display:none}.article{padding:42px 6vw}}@media(max-width:720px){.layout{display:block}.sidebar{position:fixed;z-index:30;left:0;top:48px;width:min(300px,85vw);height:calc(100vh - 48px);transform:translateX(-110%);transition:transform .22s;box-shadow:8px 0 25px #1233}.sidebar.open{transform:translateX(0)}.mobile-head{display:flex;position:sticky;z-index:40;top:0;height:48px;background:#fff;border-bottom:1px solid #dfe8ec;align-items:center;justify-content:space-between;padding:0 16px;font-size:13px;font-weight:750}.mobile-head button{border:0;background:none;color:#067185;padding:8px}.article{padding:27px 19px 40px}.article h1{font-size:28px}.lead{font-size:15px}.doc-section{grid-template-columns:30px 1fr;gap:6px}.doc-section h2{font-size:18px}.product-shot img{aspect-ratio:auto}}@media(prefers-reduced-motion:reduce){*,*:before,*:after{animation:none!important;transition:none!important;scroll-behavior:auto!important}}@media print{.sidebar,.toc,.mobile-head,.search-dialog,.page-turn{display:none!important}.layout,.main-wrap{display:block}.article{padding:0}.doc-section{break-inside:avoid}.product-shot{box-shadow:none}body{background:#fff}}'''


JS = r'''const dialog=document.getElementById('search-dialog'),input=document.getElementById('global-search'),results=document.getElementById('search-results');const index=window.EOIP_SEARCH||[];function esc(s){return String(s).replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]))}function search(){let q=input.value.trim().toLowerCase();if(!q){results.innerHTML='<div class="search-empty">Search titles, sections and instructions across every page.</div>';return}let words=q.split(/\s+/);let hits=index.map(x=>({x,score:words.reduce((n,w)=>n+(x.title.toLowerCase().includes(w)?4:0)+(x.heading.toLowerCase().includes(w)?3:0)+(x.body.toLowerCase().includes(w)?1:0),0)})).filter(v=>v.score>=words.length).sort((a,b)=>b.score-a.score).slice(0,30);results.innerHTML=hits.length?hits.map(({x})=>'<a href="'+esc(x.href)+'"><strong>'+esc(x.heading)+'</strong><small>'+esc(x.title)+'</small><p>'+esc(x.body.slice(0,175))+'...</p></a>').join(''):'<div class="search-empty">No matches. Try a shorter term or a feature name.</div>'}function openSearch(){dialog.showModal();input.value='';search();input.focus()}document.getElementById('search-open').addEventListener('click',openSearch);document.getElementById('mobile-search').addEventListener('click',openSearch);document.getElementById('search-close').addEventListener('click',()=>dialog.close());input.addEventListener('input',search);document.addEventListener('keydown',e=>{if(e.key==='/'&&!/INPUT|TEXTAREA/.test(document.activeElement.tagName)){e.preventDefault();openSearch()}if(e.key==='Escape'&&dialog.open)dialog.close()});document.getElementById('menu-toggle').addEventListener('click',()=>document.querySelector('.sidebar').classList.toggle('open'));dialog.addEventListener('click',e=>{if(e.target===dialog)dialog.close()});'''


CSS += ".downloads{border-top:1px solid #e0e9ec}.downloads>div{display:flex;justify-content:space-between;align-items:center;gap:20px;padding:15px 0;border-bottom:1px solid #e0e9ec;font-size:13px}.downloads span{display:flex;gap:15px;font-weight:750}.downloads a{white-space:nowrap}@media(max-width:520px){.downloads>div{align-items:start;flex-direction:column;gap:5px}}"
CSS += ".article{animation:pageIn .28s ease both}.doc-section{transition:background .2s ease,padding-left .2s ease}.doc-section:target{background:#eef7f8;padding-left:8px}.product-shot img{transition:filter .2s ease}.product-shot a:hover img{filter:contrast(1.035)}@keyframes pageIn{from{opacity:.45;transform:translateY(6px)}to{opacity:1;transform:translateY(0)}}@media(prefers-reduced-motion:reduce){.article{animation:none}.doc-section,.product-shot img{transition:none}}"

GUIDES = [
    ("EOIP Quick Start", "quick-start", ["connect"]),
    ("EOIP User Guide", "user-guide", ["explore", "studio", "reports"]),
    ("EOIP Administrator Guide", "administrator-guide", ["data-admin", "security"]),
    ("EOIP Operations and Troubleshooting", "operations-troubleshooting", ["troubleshoot"]),
    ("EOIP Documentation and Product Status", "documentation-status", ["reports"]),
]


def guide_sections(page, slug):
    if slug == "documentation-status":
        return [section for section in page["sections"] if section[0] != "Cognitive dashboards"]
    if slug == "user-guide" and page["slug"] == "reports":
        return [section for section in page["sections"] if section[0] == "Cognitive dashboards"]
    return page["sections"]


def build_docx(name, slug, pages):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.8)
    styles = doc.styles
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(24)
    styles["Title"].font.color.rgb = RGBColor(20, 43, 52)
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(8)
    for heading in ("Heading 1", "Heading 2"):
        styles[heading].font.name = "Aptos Display"
        styles[heading].font.color.rgb = RGBColor(0, 105, 124)
    doc.add_paragraph(name, "Title")
    doc.add_paragraph("KnowGridFlow / EOIP | 20 October 2026")
    doc.add_paragraph("This manual explains the current repository's primary web application. Authenticated journeys depend on the deployed API, graph backend, worker, policy and the reader's role. The screenshots show observed local states; they do not certify the entire workflow.")
    for page in pages:
        doc.add_heading(page["title"], level=1)
        doc.add_paragraph(page["summary"])
        doc.add_paragraph("Availability: " + page["status"])
        if page["image"] and not (slug == "user-guide" and page["slug"] == "reports"):
            picture = DASHBOARD if slug == "documentation-status" else SITE / "assets" / page["image"]
            if picture.exists():
                doc.add_picture(str(picture), width=Inches(6.5))
                doc.add_paragraph("Observed screenshot: offline project audit snapshot" if slug == "documentation-status" else "Observed screenshot: local primary web UI, not a production acceptance result")
        for heading, body, steps in guide_sections(page, slug):
            doc.add_heading(heading, level=2)
            doc.add_paragraph(body)
            for i, step in enumerate(steps, 1):
                doc.add_paragraph(f"{i}. {step}", style="Normal")
    footer = sec.footer.paragraphs[0]
    footer.text = "EOIP self help | Verify deployment behavior before relying on a feature"
    doc.save(DOCS / f"{slug}.docx")


def build_pdf(name, slug, pages):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ManualTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=21, leading=26, textColor=colors.HexColor("#173440"), spaceAfter=11))
    styles.add(ParagraphStyle(name="ManualH1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=15, leading=20, textColor=colors.HexColor("#00778e"), spaceBefore=20, spaceAfter=8))
    styles.add(ParagraphStyle(name="ManualH2", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=15, textColor=colors.HexColor("#284c59"), spaceBefore=10, spaceAfter=5))
    styles.add(ParagraphStyle(name="ManualBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.4, leading=14, spaceAfter=6))
    styles.add(ParagraphStyle(name="ManualStep", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=13.5, leftIndent=15, firstLineIndent=-13, spaceAfter=4))
    story = [Paragraph(escape(name), styles["ManualTitle"]), Paragraph("KnowGridFlow / EOIP | 20 October 2026", styles["ManualBody"]), Paragraph("This manual explains the primary web application and its documented local profiles. Authenticated behavior depends on a configured deployment and authorized role. Screenshots show observed local states, not full-product acceptance.", styles["ManualBody"])]
    for page in pages:
        story.append(Paragraph(escape(page["title"]), styles["ManualH1"]))
        story.append(Paragraph(escape(page["summary"]), styles["ManualBody"]))
        story.append(Paragraph("Availability: " + escape(page["status"]), styles["ManualBody"]))
        if page["image"] and not (slug == "user-guide" and page["slug"] == "reports"):
            picture = DASHBOARD if slug == "documentation-status" else SITE / "assets" / page["image"]
            if picture.exists():
                width = 310 if slug in {"user-guide", "documentation-status"} else 400
                img = Image(str(picture), width=width, height=width / 1.6)
                story.extend([img, Spacer(1, 8)])
        for heading, body, steps in guide_sections(page, slug):
            story.append(Paragraph(escape(heading), styles["ManualH2"]))
            story.append(Paragraph(escape(body), styles["ManualBody"]))
            for i, step in enumerate(steps, 1):
                story.append(Paragraph(f"{i}. {escape(step)}", styles["ManualStep"]))
    def decorate(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#dce8ec"))
        canvas.line(48, 42, 564, 42)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#6d8590"))
        canvas.drawString(48, 29, "EOIP self help | Verify deployment behavior")
        canvas.drawRightString(564, 29, str(document.page))
        canvas.restoreState()
    SimpleDocTemplate(str(DOCS / f"{slug}.pdf"), pagesize=letter, rightMargin=48, leftMargin=48, topMargin=48, bottomMargin=58).build(story, onFirstPage=decorate, onLaterPages=decorate)


def main():
    SITE.mkdir(parents=True, exist_ok=True)
    DOCS.mkdir(parents=True, exist_ok=True)
    assets = SITE / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    (assets / "help.css").write_text(CSS, encoding="utf-8")
    (assets / "help.js").write_text(JS, encoding="utf-8")
    if DASHBOARD.exists():
        (assets / "audit-dashboard.png").write_bytes(DASHBOARD.read_bytes())
    index = []
    for page in PAGES:
        for i, (heading, body, steps) in enumerate(page["sections"], 1):
            index.append({"href": f"{page['slug']}.html#section-{i}", "title": page["title"], "heading": heading, "body": body + " " + " ".join(steps)})
    for page in PAGES:
        (SITE / f"{page['slug']}.html").write_text(html_page(page, index), encoding="utf-8")
    listing = "# EOIP self help manuals\n\nThese files match the offline [help website](../site/index.html). The site and manuals are generated by `python scripts/build_self_help.py`. Screenshots are observed local UI states, not claims of release acceptance.\n\n"
    for name, slug, _ in GUIDES:
        listing += f"- {name}: [{slug}.pdf]({slug}.pdf) | [{slug}.docx]({slug}.docx)\n"
    (DOCS / "README.md").write_text(listing, encoding="utf-8")
    for name, slug, slugs in GUIDES:
        pages = [next(p for p in PAGES if p["slug"] == s) for s in slugs]
        build_docx(name, slug, pages)
        build_pdf(name, slug, pages)
    print(json.dumps({"site_pages": len(PAGES), "manuals": len(GUIDES), "pdfs": len(GUIDES), "docx": len(GUIDES)}))


if __name__ == "__main__":
    main()
