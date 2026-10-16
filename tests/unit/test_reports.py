"""Unit tests for the report generator, schedules and reports router (SR-NLQ-008)."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest

from eoip.reports import (
    BUILTIN_TEMPLATES,
    ReportContext,
    ReportGenerator,
    ReportScheduleStore,
    next_run_after,
    template_by_id,
)
from eoip.reports.router import build_router


def _ctx(template_id: str, **params) -> ReportContext:
    return ReportContext(
        title="Test Report",
        params=params,
        graph_stats={"vertices": 10, "edges": 42},
        evidence_manifest=[{"ref": "urn:eoip:provenance:evidence:abc"}],
        findings=["finding one", "GAP: missing control"],
        narrative="Everything is explainable.",
    )


# ---------------------------------------------------------------------------
# Templates
# ---------------------------------------------------------------------------

class TestTemplates:
    def test_five_builtin_templates(self):
        ids = {t.template_id for t in BUILTIN_TEMPLATES}
        assert ids == {
            "impact-assessment", "incident-rca", "change-risk-review",
            "compliance-evidence", "executive-brief",
        }

    def test_template_lookup(self):
        assert template_by_id("incident-rca") is not None
        assert template_by_id("nope") is None


# ---------------------------------------------------------------------------
# Generator
# ---------------------------------------------------------------------------

class TestGenerator:
    @pytest.mark.parametrize("template_id", [t.template_id for t in BUILTIN_TEMPLATES])
    @pytest.mark.parametrize("fmt", ["markdown", "html", "json"])
    def test_render_all_formats(self, template_id, fmt):
        generator = ReportGenerator()
        ctx = _ctx(template_id, subject="s", incident_id="i", proposal_id="p",
                   framework="GDPR", period="2026-Q1")
        report = generator.render(template_id, ctx, report_format=fmt, generated_at="2026-01-01")
        assert report.template == template_id
        assert report.format == fmt
        assert report.content
        assert report.report_id.startswith("urn:eoip:reports:report:")

    def test_markdown_is_self_contained(self):
        report = ReportGenerator().render("impact-assessment", _ctx("impact-assessment", subject="s"))
        assert report.content.startswith("# Test Report")
        assert "## Evidence appendix" in report.content
        assert "urn:eoip:provenance:evidence:abc" in report.content
        assert "| metric |" in report.content  # table rendered

    def test_html_is_self_contained(self):
        report = ReportGenerator().render(
            "incident-rca", _ctx("incident-rca", incident_id="i-1"), report_format="html"
        )
        assert report.content.startswith("<!doctype html>")
        assert "<style>" in report.content
        assert "<table>" in report.content or "<ul>" in report.content
        assert "</html>" in report.content

    def test_external_formats_raise(self):
        for fmt in ("pdf", "docx", "pptx"):
            with pytest.raises(ValueError, match="External gate"):
                ReportGenerator().render(
                    "executive-brief", _ctx("executive-brief"), report_format=fmt
                )

    def test_unknown_format_raises(self):
        with pytest.raises(ValueError):
            ReportGenerator().render("executive-brief", _ctx("executive-brief"), report_format="rtf")

    def test_missing_required_parameter(self):
        with pytest.raises(ValueError, match="subject"):
            ReportGenerator().render("impact-assessment", ReportContext())

    def test_render_bytes_obeys_gate_and_parameter_validation(self, monkeypatch):
        monkeypatch.delenv("EOIP_ENABLE_EXTERNAL_REPORT_FORMATS", raising=False)
        with pytest.raises(ValueError, match="External gate"):
            ReportGenerator().render_bytes("executive-brief", ReportContext())
        with pytest.raises(ValueError, match="subject"):
            ReportGenerator().render_bytes("impact-assessment", ReportContext(), allow_external_formats=True)

    def test_unknown_template(self):
        with pytest.raises(KeyError):
            ReportGenerator().render("nope", ReportContext())

    def test_evidence_manifest_preserved(self):
        report = ReportGenerator().render(
            "compliance-evidence", _ctx("compliance-evidence", framework="GDPR")
        )
        assert report.evidence_manifest == [{"ref": "urn:eoip:provenance:evidence:abc"}]


# ---------------------------------------------------------------------------
# Schedules
# ---------------------------------------------------------------------------

class TestSchedules:
    def test_create_and_due(self, tmp_path: Path):
        store = ReportScheduleStore(tmp_path / "s.json")
        schedule = store.create("executive-brief", "weekly brief", recurrence="weekly")
        assert schedule.schedule_id.startswith("urn:eoip:reports:schedule:")
        assert store.due_schedules()  # never run -> due

    def test_not_due_after_recent_run(self, tmp_path: Path):
        store = ReportScheduleStore(tmp_path / "s.json")
        schedule = store.create("executive-brief", "daily", recurrence="daily")
        store.mark_run(schedule.schedule_id, at=datetime.now(timezone.utc).isoformat())
        assert store.due_schedules() == []

    def test_due_after_period_elapsed(self, tmp_path: Path):
        store = ReportScheduleStore(tmp_path / "s.json")
        schedule = store.create("executive-brief", "daily", recurrence="daily")
        old = (datetime.now(timezone.utc) - timedelta(days=2)).isoformat()
        store.mark_run(schedule.schedule_id, at=old)
        due = store.due_schedules()
        assert [s.schedule_id for s in due] == [schedule.schedule_id]

    def test_disabled_never_due(self, tmp_path: Path):
        store = ReportScheduleStore(tmp_path / "s.json")
        schedule = store.create("executive-brief", "off", recurrence="daily")
        schedule.enabled = False
        assert store.due_schedules() == []

    def test_next_run_after(self):
        last = "2026-01-01T00:00:00+00:00"
        assert next_run_after(last, "daily").day == 2
        assert next_run_after(last, "weekly").day == 8
        assert next_run_after(last, "monthly").day == 31

    def test_persistence(self, tmp_path: Path):
        path = tmp_path / "s.json"
        store = ReportScheduleStore(path)
        schedule = store.create("incident-rca", "rca", params={"incident_id": "i"})
        reloaded = ReportScheduleStore(path)
        assert reloaded.get(schedule.schedule_id).params == {"incident_id": "i"}

    def test_delete(self, tmp_path: Path):
        store = ReportScheduleStore(tmp_path / "s.json")
        schedule = store.create("executive-brief", "x")
        assert store.delete(schedule.schedule_id)
        assert not store.delete(schedule.schedule_id)


# ---------------------------------------------------------------------------
# Router
# ---------------------------------------------------------------------------

def _client(tmp_path: Path):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    app = FastAPI()
    app.include_router(build_router(schedule_store=ReportScheduleStore(tmp_path / "s.json")))
    return TestClient(app)


class TestReportsRouter:
    def test_root_capability_endpoint(self, tmp_path: Path):
        client = _client(tmp_path)
        resp = client.get("/api/v1/reports")
        assert resp.status_code == 200
        body = resp.json()
        assert body["service"] == "reports"
        assert body["templates"] == 5

    def test_templates_endpoint(self, tmp_path: Path):
        client = _client(tmp_path)
        resp = client.get("/api/v1/reports/templates")
        assert resp.status_code == 200
        assert len(resp.json()) == 5

    def test_render_markdown(self, tmp_path: Path):
        client = _client(tmp_path)
        resp = client.post("/api/v1/reports/render", params={
            "template_id": "executive-brief", "report_format": "markdown",
        }, json={"narrative": "Summary text."})
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["content"].startswith("# ")
        assert "Summary text." in body["content"]

    def test_render_html(self, tmp_path: Path):
        client = _client(tmp_path)
        resp = client.post("/api/v1/reports/render", params={
            "template_id": "incident-rca", "report_format": "html",
        }, json={"params": {"incident_id": "i-9"}})
        assert resp.status_code == 200
        assert resp.json()["content"].startswith("<!doctype html>")

    def test_render_external_format_501(self, tmp_path: Path):
        client = _client(tmp_path)
        resp = client.post("/api/v1/reports/render", params={
            "template_id": "executive-brief", "report_format": "pdf",
        }, json={})
        assert resp.status_code == 501
        assert "External gate" in resp.json()["detail"]

    def test_render_unknown_template_404(self, tmp_path: Path):
        client = _client(tmp_path)
        resp = client.post("/api/v1/reports/render", params={"template_id": "nope"}, json={})
        assert resp.status_code == 404

    def test_schedule_endpoints(self, tmp_path: Path):
        client = _client(tmp_path)
        resp = client.post("/api/v1/reports/schedules", params={
            "template_id": "executive-brief", "name": "weekly brief",
            "recurrence": "weekly",
        })
        assert resp.status_code == 201, resp.text
        schedule_id = resp.json()["schedule_id"]
        assert client.get("/api/v1/reports/schedules").json()
        assert client.get("/api/v1/reports/schedules/due").json()
        assert client.post(f"/api/v1/reports/schedules/{schedule_id}/mark-run").status_code == 200
        assert client.delete(f"/api/v1/reports/schedules/{schedule_id}").status_code == 204
        assert client.delete(f"/api/v1/reports/schedules/{schedule_id}").status_code == 404

    def test_render_and_export_binary_formats(self, tmp_path: Path, monkeypatch):
        import io
        import zipfile

        monkeypatch.setenv("EOIP_ENABLE_EXTERNAL_REPORT_FORMATS", "true")
        client = _client(tmp_path)
        for fmt in ("pdf", "docx", "pptx"):
            resp = client.post(
                "/api/v1/reports/render",
                params={
                    "template_id": "executive-brief",
                    "report_format": fmt,
                },
                json={"narrative": f"Binary {fmt} content"},
            )
            assert resp.status_code == 200, resp.text
            report_data = resp.json()
            assert report_data["format"] == fmt
            report_id = report_data["report_id"]

            # Test export endpoint returns downloadable document
            exp_resp = client.get(f"/api/v1/reports/export/{report_id}")
            assert exp_resp.status_code == 200
            assert len(exp_resp.content) > 100
            if fmt == "pdf":
                assert exp_resp.content.startswith(b"%PDF-")
            else:
                with zipfile.ZipFile(io.BytesIO(exp_resp.content)) as archive:
                    assert archive.testzip() is None
                    assert ("word/document.xml" if fmt == "docx" else "ppt/presentation.xml") in archive.namelist()
            assert f"{report_id}.{fmt}" in exp_resp.headers["content-disposition"]

    def test_request_cannot_enable_disabled_renderer(self, tmp_path: Path, monkeypatch):
        monkeypatch.delenv("EOIP_ENABLE_EXTERNAL_REPORT_FORMATS", raising=False)
        response = _client(tmp_path).post("/api/v1/reports/render", params={
            "template_id": "executive-brief", "report_format": "pdf", "enable_external_formats": True,
        }, json={})
        assert response.status_code == 501

    def test_missing_optional_renderer_returns_503(self, tmp_path: Path, monkeypatch):
        from eoip.reports import generator

        def unavailable(*args):
            raise ImportError("renderer unavailable")

        monkeypatch.setenv("EOIP_ENABLE_EXTERNAL_REPORT_FORMATS", "true")
        monkeypatch.setattr(generator, "_sections_to_pdf", unavailable)
        response = _client(tmp_path).post("/api/v1/reports/render", params={
            "template_id": "executive-brief", "report_format": "pdf",
        }, json={})
        assert response.status_code == 503

    def test_export_format_mismatch_is_not_silently_ignored(self, tmp_path: Path):
        client = _client(tmp_path)
        response = client.post("/api/v1/reports/render", params={
            "template_id": "executive-brief", "report_format": "html",
        }, json={})
        report_id = response.json()["report_id"]
        assert client.get(f"/api/v1/reports/export/{report_id}?report_format=pdf").status_code == 409

    def test_export_unknown_report_returns_404(self, tmp_path: Path):
        client = _client(tmp_path)
        resp = client.get("/api/v1/reports/export/urn:eoip:reports:report:missing")
        assert resp.status_code == 404

    def test_report_artifact_store_persistence_and_reload(self, tmp_path: Path):
        from eoip.reports.store import ReportArtifactStore
        from eoip.reports.generator import RenderedReport

        store1 = ReportArtifactStore(tmp_path / "reports")
        rep = RenderedReport(
            report_id="urn:eoip:reports:report:rep-1",
            template="executive-brief",
            title="Durable Report",
            format="markdown",
            generated_at="2026-10-14T12:00:00Z",
            content="# Durable Content",
        )
        store1.put(rep)
        assert (tmp_path / "reports" / "index.json").exists()

        # Reload in a fresh store instance
        store2 = ReportArtifactStore(tmp_path / "reports")
        loaded = store2.get("urn:eoip:reports:report:rep-1")
        assert loaded is not None
        assert loaded.title == "Durable Report"
        assert loaded.content == "# Durable Content"
        assert len(store2.list()) == 1

        # Delete
        assert store2.delete("urn:eoip:reports:report:rep-1") is True
        store3 = ReportArtifactStore(tmp_path / "reports")
        assert store3.get("urn:eoip:reports:report:rep-1") is None

    def test_pptx_section_table_rendering(self, monkeypatch):
        import io
        import pptx
        monkeypatch.setenv("EOIP_ENABLE_EXTERNAL_REPORT_FORMATS", "true")
        ctx = _ctx("executive-brief")
        ctx.graph_stats = {"nodes": "12", "edges": "34"}
        raw_bytes = ReportGenerator().render_bytes("executive-brief", ctx, report_format="pptx", allow_external_formats=True)
        prs = pptx.Presentation(io.BytesIO(raw_bytes))
        tables = [shape.table for slide in prs.slides for shape in slide.shapes if shape.has_table]
        assert len(tables) >= 1
        table = tables[0]
        assert table.cell(0, 0).text == "metric"
        assert table.cell(0, 1).text == "value"

    def test_reports_and_schedules_cross_tenant_isolation(self, tmp_path: Path):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from starlette.requests import Request
        from eoip.reports.store import ReportArtifactStore
        from eoip.reports.schedule import ReportScheduleStore

        app = FastAPI()
        art_store = ReportArtifactStore(tmp_path / "art")
        sched_store = ReportScheduleStore(tmp_path / "sched.json")
        app.include_router(build_router(schedule_store=sched_store, artifact_store=art_store))

        # Middleware that binds tenant_id from header
        @app.middleware("http")
        async def tenant_middleware(request: Request, call_next):
            tenant = request.headers.get("X-Tenant-ID")
            if tenant:
                request.state.tenant_id = tenant
            return await call_next(request)

        client = TestClient(app)

        # 1. Tenant-alpha renders a report
        render_resp = client.post(
            "/api/v1/reports/render",
            params={"template_id": "executive-brief", "report_format": "markdown"},
            json={"narrative": "Alpha narrative."},
            headers={"X-Tenant-ID": "tenant-alpha"},
        )
        assert render_resp.status_code == 200
        report_id = render_resp.json()["report_id"]
        assert render_resp.json()["tenant_id"] == "tenant-alpha"

        # 2. Tenant-beta attempts to export tenant-alpha's report -> 403
        resp_beta = client.get(f"/api/v1/reports/export/{report_id}", headers={"X-Tenant-ID": "tenant-beta"})
        assert resp_beta.status_code == 403
        assert "Cross-tenant access forbidden" in resp_beta.json()["detail"]

        # 3. Tenant-alpha exports own report -> 200
        resp_alpha = client.get(f"/api/v1/reports/export/{report_id}", headers={"X-Tenant-ID": "tenant-alpha"})
        assert resp_alpha.status_code == 200
        assert "Alpha narrative." in resp_alpha.text

        # 4. Tenant-alpha creates a schedule
        sched_resp = client.post(
            "/api/v1/reports/schedules",
            params={"template_id": "executive-brief", "name": "Alpha Weekly Brief"},
            headers={"X-Tenant-ID": "tenant-alpha"},
        )
        assert sched_resp.status_code == 201
        schedule_id = sched_resp.json()["schedule_id"]
        assert sched_resp.json()["tenant_id"] == "tenant-alpha"

        # 5. Tenant-beta lists schedules -> does NOT see tenant-alpha's schedule
        list_beta = client.get("/api/v1/reports/schedules", headers={"X-Tenant-ID": "tenant-beta"})
        assert list_beta.status_code == 200
        assert len([s for s in list_beta.json() if s["schedule_id"] == schedule_id]) == 0

        # 6. Tenant-alpha lists schedules -> DOES see schedule
        list_alpha = client.get("/api/v1/reports/schedules", headers={"X-Tenant-ID": "tenant-alpha"})
        assert list_alpha.status_code == 200
        assert len([s for s in list_alpha.json() if s["schedule_id"] == schedule_id]) == 1

        # 7. Tenant-beta attempts to delete tenant-alpha's schedule -> 403
        del_beta = client.delete(f"/api/v1/reports/schedules/{schedule_id}", headers={"X-Tenant-ID": "tenant-beta"})
        assert del_beta.status_code == 403

        # 8. Tenant-alpha deletes own schedule -> 204
        del_alpha = client.delete(f"/api/v1/reports/schedules/{schedule_id}", headers={"X-Tenant-ID": "tenant-alpha"})
        assert del_alpha.status_code == 204

    def test_schedule_run_execution_persists_artifact(self, tmp_path: Path):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from starlette.requests import Request
        from eoip.reports.store import ReportArtifactStore
        from eoip.reports.schedule import ReportScheduleStore

        app = FastAPI()
        art_store = ReportArtifactStore(tmp_path / "art")
        sched_store = ReportScheduleStore(tmp_path / "sched.json")
        app.include_router(build_router(schedule_store=sched_store, artifact_store=art_store))

        @app.middleware("http")
        async def tenant_middleware(request: Request, call_next):
            tenant = request.headers.get("X-Tenant-ID")
            if tenant:
                request.state.tenant_id = tenant
            return await call_next(request)

        client = TestClient(app)

        # 1. Tenant-alpha creates a schedule
        sched_resp = client.post(
            "/api/v1/reports/schedules",
            params={"template_id": "executive-brief", "name": "Monthly Exec", "report_format": "markdown"},
            headers={"X-Tenant-ID": "tenant-alpha"},
        )
        assert sched_resp.status_code == 201
        sched_id = sched_resp.json()["schedule_id"]

        # 2. Tenant-beta attempts to trigger tenant-alpha's schedule -> 403
        run_beta = client.post(f"/api/v1/reports/schedules/{sched_id}/run", headers={"X-Tenant-ID": "tenant-beta"})
        assert run_beta.status_code == 403
        assert "Cross-tenant access forbidden" in run_beta.json()["detail"]

        # 3. Tenant-alpha runs their schedule -> 200
        run_alpha = client.post(f"/api/v1/reports/schedules/{sched_id}/run", headers={"X-Tenant-ID": "tenant-alpha"})
        assert run_alpha.status_code == 200
        artifact = run_alpha.json()
        assert artifact["template"] == "executive-brief"
        assert artifact["tenant_id"] == "tenant-alpha"
        assert artifact["report_id"]

        # 4. Verify artifact exists in art_store and is retrievable by tenant-alpha
        stored = art_store.get(artifact["report_id"], tenant_id="tenant-alpha")
        assert stored is not None
        assert stored.tenant_id == "tenant-alpha"

        # 5. Verify tenant-beta cannot export the generated artifact
        exp_beta = client.get(f"/api/v1/reports/export/{artifact['report_id']}", headers={"X-Tenant-ID": "tenant-beta"})
        assert exp_beta.status_code == 403



